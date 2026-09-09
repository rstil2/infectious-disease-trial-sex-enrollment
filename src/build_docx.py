#!/usr/bin/env python3
"""Build Word files for the Trials R1 package."""

import re
from pathlib import Path

import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT

_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
_ITALIC_RE = re.compile(r"\*(.+?)\*")


def set_run_font(run, size=12, bold=False, italic=False):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), "Times New Roman")
    rFonts.set(qn("w:hAnsi"), "Times New Roman")


def _add_italic_spans(p, text, size, bold):
    """Add runs for `text`, turning single-*-wrapped spans italic."""
    pos = 0
    for m in _ITALIC_RE.finditer(text):
        if m.start() > pos:
            run = p.add_run(text[pos:m.start()])
            set_run_font(run, size=size, bold=bold, italic=False)
        run = p.add_run(m.group(1))
        set_run_font(run, size=size, bold=bold, italic=True)
        pos = m.end()
    if pos < len(text):
        run = p.add_run(text[pos:])
        set_run_font(run, size=size, bold=bold, italic=False)


def add_rich_runs(p, text, size=12, base_bold=False):
    """Add runs to an existing paragraph, rendering **bold** and *italic* spans."""
    pos = 0
    for m in _BOLD_RE.finditer(text):
        if m.start() > pos:
            _add_italic_spans(p, text[pos:m.start()], size, base_bold)
        _add_italic_spans(p, m.group(1), size, True)
        pos = m.end()
    if pos < len(text):
        _add_italic_spans(p, text[pos:], size, base_bold)


def add_p(doc, text, *, bold=False, italic=False, size=12, center=False, space_after=8):
    """Add a paragraph, rendering inline **bold** and *italic* markdown spans."""
    p = doc.add_paragraph()
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    if italic:
        # Explicit italic override (e.g. table footnotes): no inline parsing needed.
        run = p.add_run(text)
        set_run_font(run, size=size, bold=bold, italic=True)
        return p
    add_rich_runs(p, text, size=size, base_bold=bold)
    return p


def heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for r in h.runs:
        set_run_font(r, size=14 if level == 1 else 12, bold=True)
    return h


def add_md_blocks(doc, path: Path):
    """Very small markdown-to-docx: headings, paragraphs, bullets."""
    lines = path.read_text().splitlines()
    buf = []

    def flush():
        nonlocal buf
        if not buf:
            return
        text = " ".join(buf).strip()
        if text:
            add_p(doc, text)
        buf = []

    for line in lines:
        if line.startswith("# "):
            flush()
            add_p(doc, line[2:].strip(), bold=True, size=16, center=True)
        elif line.startswith("## "):
            flush()
            heading(doc, line[3:].strip(), 1)
        elif line.startswith("### "):
            flush()
            heading(doc, line[4:].strip(), 2)
        elif line.startswith("---"):
            flush()
        elif line.startswith("|") or line.startswith("See `"):
            flush()
        elif line.startswith("- "):
            flush()
            p = doc.add_paragraph(style="List Bullet")
            add_rich_runs(p, line[2:].strip())
        elif line.strip() == "":
            flush()
        else:
            buf.append(line.strip())
    flush()


def add_table1(doc):
    tl = pd.read_csv(ROOT / "results" / "trial_level_estimates.csv")
    sub = tl[tl.analysis_set == "primary"].copy()
    heading(doc, "Table 1. Mixed-sex primary trials: observed female enrollment versus expected share", 1)
    headers = ["Trial", "Disease", "N", "Female n", "Observed %", "Expected %", "Ratio"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    for r in sub.itertuples(index=False):
        row = table.add_row().cells
        row[0].text = str(r.short_name)
        row[1].text = str(r.disease)
        row[2].text = f"{int(r.n):,}"
        row[3].text = f"{int(r.n_female):,}"
        row[4].text = f"{100*r.p:.1f}"
        row[5].text = f"{100*r.pi:.0f}"
        row[6].text = f"{r.ratio:.2f}"
    add_p(
        doc,
        "Expected share is defined in Methods (indication, eligibility, and target-population epidemiology). "
        "Ratio = observed / expected. Sources and extraction notes: Additional File 1.",
        italic=True,
        size=10,
    )


def build_manuscript():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(1)
        s.bottom_margin = Inches(1)
        s.left_margin = Inches(1)
        s.right_margin = Inches(1)
    add_md_blocks(doc, ROOT / "Manuscript_Trials_R1.md")
    add_table1(doc)
    heading(doc, "Figure files", 1)
    add_p(doc, "Figure 1: figures/figure1_observed_vs_expected.png")
    add_p(doc, "Figure 2: figures/figure2_enrollment_ratio.png")
    add_p(doc, "Supplementary Figure S1: figures/figure_s1_reconstruction_flow.png")
    doc.save(OUT / "Manuscript_Trials_R1.docx")


def build_simple(md_name, docx_name):
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(1)
        s.bottom_margin = Inches(1)
        s.left_margin = Inches(1)
        s.right_margin = Inches(1)
    add_md_blocks(doc, ROOT / md_name)
    doc.save(OUT / docx_name)


def build_screening_table():
    doc = Document()
    heading(doc, "Supplementary Figure S1. Reconstruction of the analysis set", 1)
    fig_path = ROOT / "figures" / "figure_s1_reconstruction_flow.png"
    if fig_path.exists():
        doc.add_picture(str(fig_path), width=Inches(6.2))
    else:
        add_p(doc, "[Figure file missing: run src/analysis.py to regenerate figures/figure_s1_reconstruction_flow.png]", italic=True)

    heading(doc, "Supplementary Table S1. Reconstruction decisions for the original 15 named studies", 1)
    log = pd.read_csv(ROOT / "data" / "screening_log.csv")
    headers = list(log.columns)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = str(h)
    for r in log.itertuples(index=False):
        cells = table.add_row().cells
        for i, val in enumerate(r):
            cells[i].text = str(val)

    heading(doc, "Supplementary Table S2. Extraction notes", 1)
    ext = pd.read_csv(ROOT / "data" / "verified_trials.csv")
    keep = ["trial_id", "short_name", "n_analysis", "n_female", "expected_female", "nct_or_registry", "doi", "extraction_source"]
    ext = ext[keep]
    table = doc.add_table(rows=1, cols=len(keep))
    table.style = "Table Grid"
    for i, h in enumerate(keep):
        table.rows[0].cells[i].text = h
    for r in ext.itertuples(index=False):
        cells = table.add_row().cells
        for i, val in enumerate(r):
            cells[i].text = str(val)

    heading(doc, "Supplementary File S3. STROBE checklist", 1)
    add_p(
        doc,
        "This observational analysis of published trial reports follows the STROBE statement "
        "(von Elm et al., Ann Intern Med 2007; doi:10.7326/0003-4819-147-8-200710160-00010).",
        italic=True,
    )
    strobe = pd.read_csv(ROOT / "data" / "STROBE_checklist.csv")
    headers = list(strobe.columns)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = str(h)
    for r in strobe.itertuples(index=False):
        cells = table.add_row().cells
        for i, val in enumerate(r):
            cells[i].text = str(val)

    heading(doc, "Model specification", 1)
    add_p(
        doc,
        "Observational unit: trial. Linear predictor: logit(p) − logit(π) = α + δ_disease + residual. "
        "No calendar-time term in the primary model. Priors: α ~ Normal(0,1); τ, σδ ~ HalfNormal(0.5). "
        "Componentwise Metropolis–Hastings, 20,000 iterations, 5,000 discarded as burn-in, seed 42. "
        "Acceptance rates: α 86%, δ 58%, τ 85%, σδ 93%. "
        "See src/analysis.py and results/summary.json.",
    )
    doc.save(OUT / "Supplementary_Appendix_R1.docx")


if __name__ == "__main__":
    build_manuscript()
    build_simple("Point_by_Point_Response.md", "Point_by_Point_Response.docx")
    build_simple("Cover_Letter_Trials_R1.md", "Cover_Letter_Trials_R1.docx")
    build_screening_table()
    print("Wrote Word files to", OUT)
